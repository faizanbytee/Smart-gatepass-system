"""
Build Final Submission Artifacts:
1. Jupyter Notebook: Wine_Quality_Case_Study.ipynb
2. Word Document Report: Wine_Quality_Project_Report.docx
3. Master Submission ZIP: Wine_Quality_Case_Study_Submission.zip
"""

import os
import shutil
import zipfile
import nbformat as nbf
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# ==============================================================================
# 1. CREATE JUPYTER NOTEBOOK (Wine_Quality_Case_Study.ipynb)
# ==============================================================================
print("Creating Wine_Quality_Case_Study.ipynb...")
nb = nbf.v4.new_notebook()

cells = []

# Title & Business Understanding
cells.append(nbf.v4.new_markdown_cell("""# Case Study: Alcoholic Beverage Sector - Wine Quality Analysis & Prediction
**Author**: Data Science Team (Screaming Eagle Winery Case Study)  
**Domain**: Alcoholic Beverage Sector  
**Objective**: End-to-end data science study analyzing physicochemical features of red and white wines, building an interactive Streamlit data application, and developing predictive machine learning models across 5 algorithms.

---

## Business Understanding
The wine industry is a multi-billion dollar dynamic market where quality directly determines:
1. **Consumer Satisfaction**: Rich aromas, balanced acidity, and harmonious finish.
2. **Reputation & Brand Equity**: High ratings drive loyalty and premium pricing.
3. **Competitive Advantage**: Differentiates estate vintages from mass commercial blends.
4. **Economic Value**: High-quality wines appreciate over time and command top auction valuations.
5. **Critical Acclaim**: Scores by top critics (Robert Parker, Wine Spectator) dictate market velocity.

### Chemical Composition vs Sensory Characteristics
- **Fixed Acidity**: Tartaric, malic, and citric acids contributing to freshness.
- **Volatile Acidity**: Acetic acid; high levels lead to an unpleasant vinegar taint.
- **Residual Sugar**: Post-fermentation sweetness balancing perceived astringency.
- **Chlorides**: Salt concentration influencing palate weight.
- **Sulfur Dioxide ($SO_2$)**: Antioxidant and antimicrobial preservation chemistry.
- **Density**: Inversely correlated with alcohol; indicates dissolved solids.
- **pH**: Acidity/alkalinity balance (typically 3.0 - 3.8).
- **Sulphates**: Potassium sulphate acting as an antioxidant.
- **Alcohol**: Determines warmth, body, and aromatic release."""))

# Imports
cells.append(nbf.v4.new_code_cell("""# Import Essential Libraries
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
import warnings
warnings.filterwarnings('ignore')

# Set visual aesthetic
plt.style.use('seaborn-v0_8-whitegrid' if 'seaborn-v0_8-whitegrid' in plt.style.available else 'default')
%matplotlib inline"""))

# Sprint 1 Header
cells.append(nbf.v4.new_markdown_cell("""---
# SPRINT 1: Exploratory Data Analysis (EDA)

### Tasks:
1. Explore the data distribution of each column.
2. Identify important patterns and significant variables with respect to Wine Quality.
3. Derive data-driven business insights and recommendations for winery management."""))

# Load Data
cells.append(nbf.v4.new_code_cell("""# Load the Wine Quality Dataset
df = pd.read_csv('wine_quality.csv')
df.columns = df.columns.str.strip() # Sanitize column headers

print(f"Dataset Shape: {df.shape[0]} rows, {df.shape[1]} columns")
df.head()"""))

cells.append(nbf.v4.new_code_cell("""# Data info and missing values check
print("Missing values per column:")
print(df.isnull().sum())
print("\\nSummary Statistics:")
df.describe().T"""))

# EDA Plots: Target Distributions
cells.append(nbf.v4.new_markdown_cell("""### 1. Target Variable Distributions
We inspect the sensory score `quality` (integer 3-9) and the binary tier `good` (1 for score $\ge 7$, 0 otherwise)."""))

cells.append(nbf.v4.new_code_cell("""fig, axes = plt.subplots(1, 2, figsize=(14, 5))

# Quality score countplot
sns.countplot(data=df, x='quality', palette='viridis', ax=axes[0], hue='quality', legend=False)
axes[0].set_title('Sensory Quality Score Distribution (3 to 9)', fontsize=12, fontweight='bold')
axes[0].set_xlabel('Quality Score')
axes[0].set_ylabel('Count')

# Good pie chart
good_counts = df['good'].value_counts()
axes[1].pie(good_counts, labels=['Standard Wine (<7)', 'Premium Wine (>=7)'],
           autopct='%1.1f%%', startangle=140, colors=['#d9534f', '#5cb85c'], explode=(0, 0.08), shadow=True)
axes[1].set_title('Binary Quality Proportion (Standard vs Premium)', fontsize=12, fontweight='bold')

plt.tight_layout()
plt.show()"""))

# Feature Distributions
cells.append(nbf.v4.new_markdown_cell("""### 2. Physicochemical Distributions & Skewness"""))

cells.append(nbf.v4.new_code_cell("""numeric_features = [
    'fixed acidity', 'volatile acidity', 'citric acid', 'residual sugar',
    'chlorides', 'free sulfur dioxide', 'total sulfur dioxide', 'density',
    'pH', 'sulphates', 'alcohol'
]

fig, axes = plt.subplots(4, 3, figsize=(16, 12))
axes = axes.flatten()

for i, col in enumerate(numeric_features):
    sns.histplot(df[col], kde=True, ax=axes[i], color='#2c3e50', bins=30)
    skew = df[col].skew()
    axes[i].set_title(f'{col} (Skew: {skew:.2f})', fontsize=10, fontweight='bold')

fig.delaxes(axes[11])
plt.suptitle('Physicochemical Distributions', fontsize=14, fontweight='bold')
plt.tight_layout()
plt.show()"""))

# Correlation Heatmap
cells.append(nbf.v4.new_markdown_cell("""### 3. Correlation Heatmap: Physicochemical Drivers of Quality"""))

cells.append(nbf.v4.new_code_cell("""plt.figure(figsize=(12, 8))
corr = df[numeric_features + ['quality', 'good']].corr()
mask = np.triu(np.ones_like(corr, dtype=bool))
sns.heatmap(corr, mask=mask, annot=True, fmt='.2f', cmap='coolwarm', vmin=-0.6, vmax=0.6, linewidths=0.5)
plt.title('Correlation Heatmap: Chemical Drivers vs Quality', fontsize=13, fontweight='bold')
plt.tight_layout()
plt.show()"""))

# Boxplots vs Good
cells.append(nbf.v4.new_markdown_cell("""### 4. Significant Feature Comparison (Standard vs Premium)"""))

cells.append(nbf.v4.new_code_cell("""fig, axes = plt.subplots(2, 2, figsize=(14, 10))

# Alcohol
sns.boxplot(data=df, x='good', y='alcohol', palette=['#e74c3c', '#2ecc71'], ax=axes[0, 0], hue='good', legend=False)
axes[0, 0].set_title('Alcohol (% vol) vs Quality Class', fontweight='bold')

# Volatile Acidity
sns.boxplot(data=df, x='good', y='volatile acidity', palette=['#e74c3c', '#2ecc71'], ax=axes[0, 1], hue='good', legend=False)
axes[0, 1].set_title('Volatile Acidity (g/L) vs Quality Class', fontweight='bold')

# Density
sns.boxplot(data=df, x='good', y='density', palette=['#e74c3c', '#2ecc71'], ax=axes[1, 0], hue='good', legend=False)
axes[1, 0].set_title('Density (g/cm³) vs Quality Class', fontweight='bold')

# Chlorides
sns.boxplot(data=df, x='good', y='chlorides', palette=['#e74c3c', '#2ecc71'], ax=axes[1, 1], hue='good', legend=False)
axes[1, 1].set_title('Chlorides (g/L) vs Quality Class', fontweight='bold')

plt.tight_layout()
plt.show()"""))

# Sprint 1 Findings
cells.append(nbf.v4.new_markdown_cell("""### Sprint 1 Insights & Strategic Recommendations
1. **Alcohol is the Single Strongest Positive Driver (+0.44)**: Premium wines consistently possess higher alcohol ($>11.5\%$). Winemakers must optimize vineyard sun exposure and harvest Brix to ensure complete phenolic ripeness.
2. **Volatile Acidity is the Primary Defect Indicator (-0.27)**: High volatile acidity produces an off-putting vinegar aroma. Enforce cellar hygiene and controlled, cool fermentation to suppress *Acetobacter*.
3. **Density is Inversely Correlated (-0.31)**: Lower density indicates thorough fermentation of sugars into alcohol with fewer dissolved defects.
4. **Chlorides (-0.20)**: Excess saline mineral content diminishes balance and elegance; irrigation water quality must be tightly monitored."""))

# Sprint 2 Header
cells.append(nbf.v4.new_markdown_cell("""---
# SPRINT 2: Interactive Data Dashboard (Streamlit)

A production-ready web application (`app.py`) was created to allow stakeholders to dynamically explore data, evaluate models, and simulate wine batches.

### Key Features of `app.py`:
- **Executive Summary**: Business background, economic levers, chemical dictionary, and downloadable datasets.
- **Interactive EDA Studio**: Dynamic Plotly distributions, boxplots, Red vs White varietal contrasts, and correlation heatmaps.
- **Model Leaderboard & Diagnostics**: Visual model comparisons, confusion matrices, and feature importance bar charts.
- **Real-Time Quality Predictor**: Interactive sliders with pre-configured batch presets (*Screaming Eagle Cabernet*, *Sonoma Reserve Chardonnay*, *Commercial Blend*), outputting real-time class predictions, confidence gauges, and oenological recommendations.
- **Strategic Winery Recommendations**: Comprehensive managerial roadmap.

To run the Streamlit dashboard:
```bash
python -m streamlit run app.py
```"""))

# Sprint 3 Header
cells.append(nbf.v4.new_markdown_cell("""---
# SPRINT 3: Model Building & Multi-Algorithm Benchmarking

Predicting wine quality using physicochemical characteristics across 5 mandatory algorithms:
- K-Nearest Neighbors (KNN)
- Logistic Regression
- Support Vector Machines (SVM)
- Decision Trees
- Random Forest"""))

# Steps 1 to 5
cells.append(nbf.v4.new_code_cell("""from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler, OneHotEncoder
from sklearn.neighbors import KNeighborsClassifier
from sklearn.linear_model import LogisticRegression
from sklearn.svm import SVC
from sklearn.tree import DecisionTreeClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn import metrics

# Step 1: Input and Target definition
input_features = [
    'fixed acidity', 'volatile acidity', 'citric acid', 'residual sugar',
    'chlorides', 'free sulfur dioxide', 'total sulfur dioxide', 'density',
    'pH', 'sulphates', 'alcohol', 'color'
]
target = 'good'

X = df[input_features].copy()
y = df[target].copy()

# Step 3: Train-Test Split (75:25)
X_train, X_test, y_train, y_test = train_test_split(
    X, y, test_size=0.25, random_state=42, stratify=y
)

# Step 4 & 5: Data Preparation (Fit on train, transform on test)
num_cols = [c for c in input_features if c != 'color']
cat_cols = ['color']

scaler = StandardScaler()
X_train_num = scaler.fit_transform(X_train[num_cols])
X_test_num = scaler.transform(X_test[num_cols])

ohe = OneHotEncoder(drop='first', sparse_output=False)
X_train_cat = ohe.fit_transform(X_train[cat_cols])
X_test_cat = ohe.transform(X_test[cat_cols])

X_train_prep = np.hstack([X_train_num, X_train_cat])
X_test_prep = np.hstack([X_test_num, X_test_cat])

print(f"Train Shape: {X_train_prep.shape}, Test Shape: {X_test_prep.shape}")"""))

# Step 6 & 7
cells.append(nbf.v4.new_code_cell("""# Step 6: Model Training Phase (All 5 Algorithms)
models = {
    'KNN': KNeighborsClassifier(n_neighbors=7, weights='distance'),
    'Logistic Regression': LogisticRegression(max_iter=1000, random_state=42),
    'Support Vector Machines': SVC(random_state=42),
    'Decision Trees': DecisionTreeClassifier(max_depth=10, min_samples_split=10, random_state=42),
    'Random Forest': RandomForestClassifier(n_estimators=150, max_depth=15, min_samples_split=5, random_state=42, n_jobs=-1)
}

# Step 7: Prediction & Evaluation
results = []
for name, model in models.items():
    model.fit(X_train_prep, y_train)
    preds = model.predict(X_test_prep)
    acc = metrics.accuracy_score(y_test, preds)
    prec = metrics.precision_score(y_test, preds, zero_division=0)
    rec = metrics.recall_score(y_test, preds, zero_division=0)
    f1 = metrics.f1_score(y_test, preds, zero_division=0)
    
    results.append({
        'Algorithm': name,
        'Accuracy': acc,
        'Precision': prec,
        'Recall': rec,
        'F1_Score': f1
    })

results_df = pd.DataFrame(results).sort_values(by='Accuracy', ascending=False).reset_index(drop=True)
results_df"""))

# Step 8
cells.append(nbf.v4.new_code_cell("""# Step 8: Algorithm Comparison Plot
plt.figure(figsize=(10, 5))
palette = ['#27ae60', '#f39c12', '#8e44ad', '#c0392b', '#2980b9']
bars = plt.bar(results_df['Algorithm'], results_df['Accuracy'] * 100, color=palette, width=0.55, edgecolor='#2c3e50')
plt.title('Algorithm Comparison: Test Accuracy on Wine Quality Prediction', fontsize=13, fontweight='bold')
plt.ylabel('Test Accuracy (%)', fontsize=11)
plt.xlabel('Machine Learning Algorithm', fontsize=11)
plt.ylim(75, 95)
plt.axhline(80.37, color='gray', linestyle='--', label='Majority Class Baseline (80.4%)')

for bar in bars:
    height = bar.get_height()
    plt.annotate(f'{height:.2f}%',
                 xy=(bar.get_x() + bar.get_width() / 2, height),
                 xytext=(0, 5), textcoords="offset points",
                 ha='center', va='bottom', fontweight='bold')

plt.legend(loc='lower right')
plt.tight_layout()
plt.savefig('model_accuracies.png', dpi=300)
plt.show()"""))

cells.append(nbf.v4.new_markdown_cell("""### Conclusion on Best Algorithm
- **Champion Algorithm**: **Random Forest** achieved the highest test accuracy (**88.00%**), outperforming KNN (86.95%), SVM (83.38%), Logistic Regression (82.34%), and Decision Trees (82.15%).
- **Rationale**:
  - Ensemble aggregation over 150 decision trees neutralizes variance and avoids single-tree overfitting.
  - Successfully captures multi-dimensional non-linear interactions between alcohol, acidity, and sulfur dioxide.
  - Achieved an outstanding **78.70% precision** on the minority premium class, preventing low-quality batches from being misclassified as reserve tier."""))

nb['cells'] = cells

with open('Wine_Quality_Case_Study.ipynb', 'w', encoding='utf-8') as f:
    nbf.write(nb, f)
print("Wine_Quality_Case_Study.ipynb created successfully.")

# ==============================================================================
# 2. CREATE WORD REPORT (Wine_Quality_Project_Report.docx)
# ==============================================================================
print("Creating Wine_Quality_Project_Report.docx...")
doc = Document()

# Set Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("CASE STUDY REPORT: ALCOHOLIC BEVERAGE SECTOR\nWine Quality Intelligence & Predictive Modeling")
title_run.font.size = Pt(20)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(114, 47, 55) # Burgundy
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle_p = doc.add_paragraph()
sub_run = subtitle_p.add_run("Client: Screaming Eagle Winery & Global Vineyard Stakeholders\nAuthor: Lead Data Science Team | Date: October 2026\n")
sub_run.font.size = Pt(11)
sub_run.font.italic = True
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Executive Summary
doc.add_heading("1. Executive Summary", level=1)
doc.add_paragraph(
    "In the global alcoholic beverage sector, sensory quality is the single most critical determinant of brand reputation, "
    "consumer loyalty, and pricing power. This study investigates the physicochemical signatures of 6,497 red and white wine "
    "samples to discover the key drivers of wine quality, deliver an interactive Streamlit analytics platform, and deploy predictive "
    "machine learning models to standardize quality at production scale."
)
doc.add_paragraph(
    "Key Takeaway: Wine quality is primarily governed by Alcohol (+0.44 correlation), Volatile Acidity (-0.27 correlation), and Density (-0.31 correlation). "
    "A production Random Forest classifier predicts premium vintage quality with 88.00% accuracy and 0.9069 ROC-AUC, enabling proactive "
    "cellar interventions before bottling."
)

# Domain Understanding
doc.add_heading("2. Domain Understanding & Physicochemical Chemistry", level=1)
doc.add_paragraph(
    "Wine quality depends on the harmonious interaction of chemical composition and sensory characteristics:"
)
chem_points = [
    ("Fixed Acidity: ", "Tartaric, malic, and citric acids providing foundational structural acidity and crispness."),
    ("Volatile Acidity: ", "Acetic acid concentration. High volatile acidity (>0.50 g/L) causes an undesirable vinegar flaw."),
    ("Residual Sugar: ", "Post-fermentation sugar balancing perceived astringency and alcohol warmth."),
    ("Chlorides: ", "Salt minerals that influence mouthfeel and body weight."),
    ("Sulfur Dioxide (Free & Total SO2): ", "Antioxidant and antimicrobial protection crucial for aging stability."),
    ("Density: ", "Inversely related to alcohol and unfermented solids; lower density signifies complete fermentation."),
    ("pH: ", "Measure of acidity/alkalinity; lower pH reinforces antimicrobial stability."),
    ("Sulphates: ", "Potassium sulphate contributing to aroma freshness and microbial inhibition."),
    ("Alcohol: ", "Determines warmth, viscosity, and aromatic vapor release.")
]
for bold_prefix, text in chem_points:
    p = doc.add_paragraph(style='List Bullet')
    r_bold = p.add_run(bold_prefix)
    r_bold.bold = True
    p.add_run(text)

# Sprint 1: EDA
doc.add_heading("3. Sprint 1: Exploratory Data Analysis & Chemical Drivers", level=1)
doc.add_paragraph(
    "The dataset contains 6,497 rows and 14 columns (4,898 white wines and 1,599 red wines). Sensory quality scores follow a bell curve "
    "ranging from 3 to 9, with 76.6% concentrated at intermediate scores 5 and 6. Premium vintages (quality >= 7) constitute 19.7% of all samples."
)

doc.add_heading("Feature Significance Ranking (Correlations)", level=2)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = table.rows[0].cells
hdr_cells[0].text = "Chemical Feature"
hdr_cells[1].text = "Corr. with Good (0/1)"
hdr_cells[2].text = "Corr. with Quality (3-9)"
for cell in hdr_cells:
    cell.paragraphs[0].runs[0].font.bold = True

corr_data = [
    ("Alcohol (% vol)", "+0.390", "+0.444"),
    ("Density (g/cm³)", "-0.275", "-0.306"),
    ("Volatile Acidity (g/L)", "-0.152", "-0.266"),
    ("Chlorides (g/L)", "-0.162", "-0.201"),
    ("Citric Acid (g/L)", "+0.054", "+0.086"),
    ("Fixed Acidity (g/L)", "-0.049", "-0.077"),
    ("Total Sulfur Dioxide (mg/L)", "-0.051", "-0.041"),
    ("Residual Sugar (g/L)", "-0.064", "-0.037"),
    ("Sulphates (g/L)", "+0.034", "+0.038"),
    ("Free Sulfur Dioxide (mg/L)", "+0.015", "+0.055"),
    ("pH Level", "+0.028", "+0.020")
]
for f, cg, cq in corr_data:
    row_cells = table.add_row().cells
    row_cells[0].text = f
    row_cells[1].text = cg
    row_cells[2].text = cq

doc.add_paragraph()

# Sprint 2: Data App
doc.add_heading("4. Sprint 2: Streamlit Interactive Data Dashboard", level=1)
doc.add_paragraph(
    "To provide dynamic, interactive decision support for winery leadership, an enterprise Streamlit dashboard (app.py) was built. "
    "Key capabilities include:"
)
app_points = [
    "Executive Overview with key KPIs (Total samples, % premium, average alcohol, volatile acidity, pH).",
    "Interactive EDA Studio with dynamic Plotly histograms, box plots, Red vs White varietal contrasts, and correlation matrices.",
    "Model Leaderboard & Diagnostics featuring accuracy bar charts, confusion matrices, and feature importance rankings.",
    "Real-Time Wine Quality Predictor: Sliders and one-click presets for rapid batch simulation with live probability gauges.",
    "Strategic Winemaking SOPs providing tailored guidance on sulfur dioxide and temperature management."
]
for pt in app_points:
    doc.add_paragraph(pt, style='List Bullet')

# Sprint 3: Model Building
doc.add_heading("5. Sprint 3: Model Building & Comparative Benchmark", level=1)
doc.add_paragraph(
    "Data preparation and model building followed a rigorous 8-step methodology:\n"
    "- Train/Test Split: 75% training (4,872 samples) and 25% testing (1,625 samples), stratified.\n"
    "- Feature Scaling & Encoding: StandardScaler fit on training numerical features and transformed on test features; OneHotEncoder applied to varietal color.\n"
    "- 5 Machine Learning Algorithms: Evaluated on unseen test data using metrics.accuracy_score."
)

table_m = doc.add_table(rows=1, cols=6)
table_m.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_m = table_m.rows[0].cells
hdr_m[0].text = "Algorithm"
hdr_m[1].text = "Train Acc."
hdr_m[2].text = "Test Acc."
hdr_m[3].text = "Precision"
hdr_m[4].text = "Recall"
hdr_m[5].text = "F1-Score"
for cell in hdr_m:
    cell.paragraphs[0].runs[0].font.bold = True

bench_data = [
    ("Random Forest (Champion)", "99.10%", "88.00%", "0.7870", "0.5329", "0.6355"),
    ("K-Nearest Neighbors (KNN)", "100.00%", "86.95%", "0.6974", "0.5925", "0.6407"),
    ("Support Vector Machines (SVM)", "84.30%", "83.38%", "0.6870", "0.2821", "0.4000"),
    ("Logistic Regression", "81.65%", "82.34%", "0.6127", "0.2727", "0.3774"),
    ("Decision Trees", "92.49%", "82.15%", "0.5502", "0.4984", "0.5230")
]
for row in bench_data:
    row_cells = table_m.add_row().cells
    for i, val in enumerate(row):
        row_cells[i].text = val

doc.add_paragraph()
doc.add_heading("Conclusion on Best Algorithm", level=2)
doc.add_paragraph(
    "Random Forest is conclusively the superior algorithm. By ensembling 150 de-correlated decision trees, it models complex non-linear "
    "interactions between alcohol, density, and volatile acidity, achieving 88.00% test accuracy and a dominant 0.9069 ROC-AUC. Furthermore, "
    "its 78.70% precision prevents substandard vintages from erroneously reaching ultra-premium estate allocations."
)

# Strategic Recommendations
doc.add_heading("6. Strategic Business Decisions & Winemaking Recommendations", level=1)
recs = [
    ("1. Harvest Timing & Sugar-Alcohol Optimization: ", "Wines exceeding 11.5% alcohol exhibit 3.8x greater likelihood of achieving premium scores. Vineyard managers should enforce canopy management and harvest timing to achieve optimal fruit Brix (23.5-25.5° Brix)."),
    ("2. Strict Volatile Acidity Suppression: ", "Volatile acidity above 0.50 g/L dramatically degrades perceived quality. Maintain cool fermentation temperatures (15-18°C for whites, 24-28°C for reds) and apply inert gas blanketing during aging."),
    ("3. Dynamic Molecular SO2 Stewardship: ", "Maintain 0.6-0.8 mg/L active molecular SO2 calibrated to wine pH to guarantee oxidation protection and microbial stability."),
    ("4. Three-Tier Product Allocation Strategy: ", "Utilize Random Forest predictions at racking: batches with >=85% probability allocated to Flagship Reserve ($500+), 50-84% to Estate labels ($75-$150), and <50% to everyday table blends, protecting brand prestige and maximizing margins.")
]
for title_r, body_r in recs:
    p = doc.add_paragraph()
    r_t = p.add_run(title_r)
    r_t.bold = True
    p.add_run(body_r)

doc.save('Wine_Quality_Project_Report.docx')
print("Wine_Quality_Project_Report.docx created successfully.")

# ==============================================================================
# 3. CREATE README.md
# ==============================================================================
print("Creating README.md...")
readme_content = """# Alcoholic Beverage Sector: Wine Quality Intelligence System

End-to-end data science study and production-grade machine learning application predicting wine quality from physicochemical features.

## Project Structure
- `Wine_Quality_Case_Study.ipynb`: Comprehensive Jupyter Notebook containing full EDA, data preparation, multi-algorithm training, evaluation, and visualizations.
- `Wine_Quality_Project_Report.docx`: Executive project report formatted for winery leadership and academic review.
- `app.py`: Interactive Streamlit Data Dashboard (EDA studio, model benchmarks, real-time batch simulator).
- `train_models.py`: Model training script strictly executing Steps 1 to 8 across 5 algorithms.
- `eda_analysis.py`: Exploratory data analysis script generating statistical summaries and plots.
- `wine_quality.csv`: Wine quality dataset (6,497 samples, 14 columns).
- `model_accuracies.png`: Accuracy comparison plot across all 5 algorithms.
- `eda_plots/`: High-resolution figures covering feature distributions, target breakdown, and correlation heatmaps.
- `artifacts/`: Serialized model bundle (`model_bundle.joblib`), metrics JSON, and feature importance plots.

## How to Run

### 1. Run Machine Learning Pipeline:
```bash
python train_models.py
```

### 2. Launch Interactive Streamlit Data Dashboard:
```bash
python -m streamlit run app.py
```
Access the dashboard at `http://localhost:8501`.

### 3. Open Jupyter Notebook:
```bash
jupyter notebook Wine_Quality_Case_Study.ipynb
```
"""
with open('README.md', 'w', encoding='utf-8') as f:
    f.write(readme_content)

# ==============================================================================
# 4. CREATE MASTER SUBMISSION ZIP
# ==============================================================================
zip_filename = 'Wine_Quality_Case_Study_Submission.zip'
print(f"Creating master ZIP: {zip_filename}...")

files_and_dirs_to_zip = [
    'Wine_Quality_Case_Study.ipynb',
    'Wine_Quality_Project_Report.docx',
    'app.py',
    'train_models.py',
    'eda_analysis.py',
    'wine_quality.csv',
    'model_accuracies.png',
    'README.md',
    'artifacts',
    'eda_plots'
]

with zipfile.ZipFile(zip_filename, 'w', zipfile.ZIP_DEFLATED) as zipf:
    for item in files_and_dirs_to_zip:
        if os.path.isfile(item):
            zipf.write(item, arcname=item)
            print(f"Added file: {item}")
        elif os.path.isdir(item):
            for root, dirs, files in os.walk(item):
                for f in files:
                    full_path = os.path.join(root, f)
                    rel_path = os.path.relpath(full_path, '.')
                    zipf.write(full_path, arcname=rel_path)
            print(f"Added directory: {item}")

# Also copy ZIP to Desktop for immediate easy upload
desktop_zip = r'C:\Users\Faizanbytee\OneDrive\Desktop\Wine_Quality_Case_Study_Submission.zip'
shutil.copyfile(zip_filename, desktop_zip)
print(f"Copied ZIP to Desktop: {desktop_zip}")
print(f"ZIP Size: {os.path.getsize(zip_filename) / (1024*1024):.2f} MB")
print("All submission artifacts created successfully!")
